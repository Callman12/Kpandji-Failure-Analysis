import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen format
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Color palette
    BG_COLOR = RGBColor(0x12, 0x12, 0x12)       # Dark steel / #121212
    CARD_BG = RGBColor(0x2B, 0x2B, 0x2B)        # Lighter steel / #2B2B2B
    EMERALD = RGBColor(0x0B, 0x66, 0x23)        # Emerald Green / #0B6623
    TERRACOTTA = RGBColor(0xC8, 0x5A, 0x17)     # Ocre / Terracotta / #C85A17
    TEXT_LIGHT = RGBColor(0xEC, 0xEC, 0xEC)     # Off-white text
    TEXT_MUTED = RGBColor(0xBC, 0xBC, 0xBC)     # Muted text
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)

    blank_slide_layout = prs.slide_layouts[6] # Blank layout

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR

    def add_header(slide, title_text, category_text="KPANDJI FAILURE ANALYSIS | CASE STUDY"):
        # Category / Tracker
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = TERRACOTTA
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.9))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_LIGHT

        # Accent line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.75), Inches(2.0), Inches(0.06))
        line.fill.solid()
        line.fill.fore_color.rgb = EMERALD
        line.line.color.rgb = EMERALD

    def add_card(slide, left, top, width, height, title, points):
        # Card background shape
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = EMERALD
        shape.line.width = Pt(1.5)

        # Text box inside card
        tb = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.3), width - Inches(0.6), height - Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True
        
        # Card Title
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = EMERALD
        p.space_after = Pt(12)

        # Points
        for pt in points:
            p2 = tf.add_paragraph()
            p2.text = "• " + pt
            p2.font.size = Pt(13)
            p2.font.color.rgb = TEXT_LIGHT
            p2.space_after = Pt(8)

    # ==========================================
    # SLIDE 1: Title Slide (Dark & Impactful)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide1)

    # Decorative accent bar on left
    bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.2), Inches(0.15), Inches(3.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TERRACOTTA
    bar.line.fill.background()

    tb1 = slide1.shapes.add_textbox(Inches(1.4), Inches(2.0), Inches(10.5), Inches(3.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "Dreaming the Ivorian Automotive Industry"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    p.space_after = Pt(8)

    p_sub = tf1.add_paragraph()
    p_sub.text = "The Rise, Fall, and Lessons of Kpandji"
    p_sub.font.size = Pt(22)
    p_sub.font.color.rgb = EMERALD
    p_sub.space_after = Pt(24)

    p_meta = tf1.add_paragraph()
    p_meta.text = "A Strategic Post-Mortem Analysis in Software & Industrial Engineering\nPresented by Elie Pethuel Ehoussou | IIT L3 Software Engineering"
    p_meta.font.size = Pt(13)
    p_meta.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 2: Context & Genesis
    # ==========================================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide2)
    add_header(slide2, "1. Genesis & Ambition: The Birth of Kpandji")
    add_card(slide2, Inches(0.8), Inches(2.1), Inches(5.6), Inches(4.8), 
             "The Visionary Spark", 
             [
                 "An ambitious initiative born in Côte d'Ivoire to create a homegrown automotive brand.",
                 "Symbolized national pride, technological emancipation, and industrial independence.",
                 "Aimed to challenge foreign market dominance and showcase local engineering talent."
             ])
    add_card(slide2, Inches(6.8), Inches(2.1), Inches(5.7), Inches(4.8), 
             "The Core Narrative", 
             [
                 "« Dreaming of a national automotive industry is an act of faith; building it without an ecosystem is a leap into the void. »",
                 "Initiated by visionary entrepreneurs seeking to establish West Africa's industrial footprint.",
                 "Faced immediate structural friction between national ambition and local manufacturing realities."
             ])

    # ==========================================
    # SLIDE 3: Vision & Objectives
    # ==========================================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide3)
    add_header(slide3, "2. Strategic Objectives & Economic Goals")
    add_card(slide3, Inches(0.8), Inches(2.1), Inches(5.6), Inches(4.8), 
             "Industrial Goals", 
             [
                 "Design vehicles adapted to African roads, climate, and purchasing power.",
                 "Create local employment across assembly, maintenance, and distribution chains.",
                 "Stimulate local technological transfer and mechanical engineering know-how."
             ])
    add_card(slide3, Inches(6.8), Inches(2.1), Inches(5.7), Inches(4.8), 
             "Socio-Economic Impact", 
             [
                 "Reduce massive dependency on imported used vehicles ('Occasions d'Europe').",
                 "Position Abidjan as a regional hub for light vehicle manufacturing.",
                 "Foster public-private partnerships to back heavy industrial investments."
             ])

    # ==========================================
    # SLIDE 4: Business Model & Value Chain
    # ==========================================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide4)
    add_header(slide4, "3. Business Model: SKD / CKD Assembly Strategy")
    add_card(slide4, Inches(0.8), Inches(2.1), Inches(5.6), Inches(4.8), 
             "Semi-Knocked Down (SKD) Model", 
             [
                 "Importation of pre-assembled chassis, engines, and major sub-assemblies from Asia.",
                 "Local final assembly, bodywork customization, and interior finishing in Côte d'Ivoire.",
                 "Pragmatic starting point to bypass lack of heavy steel foundry infrastructure."
             ])
    add_card(slide4, Inches(6.8), Inches(2.1), Inches(5.7), Inches(4.8), 
             "The 'Local Car' Debate", 
             [
                 "Is it truly an Ivorian car or foreign kit assembly? A vital question of value chain definition.",
                 "In early industrialization stages, 100% local production is economically unviable.",
                 "The challenge lies in progressively increasing local component integration over time."
             ])

    # ==========================================
    # SLIDE 5: Prototype & Public Reception
    # ==========================================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide5)
    add_header(slide5, "4. The Prototype & Public Enthusiasm")
    add_card(slide5, Inches(0.8), Inches(2.1), Inches(5.6), Inches(4.8), 
             "Public Unveiling", 
             [
                 "Strong emotional resonance and national media coverage upon prototype launch.",
                 "Generational pride: seeing an Ivorian badge on a vehicle grille.",
                 "High expectations from consumers and state officials eager for industrial success."
             ])
    add_card(slide5, Inches(6.8), Inches(2.1), Inches(5.7), Inches(4.8), 
             "Early Warning Signals", 
             [
                 "Enthusiasm masked critical technical and financial vulnerabilities.",
                 "Discrepancy between lab prototypes and robust mass-market requirements.",
                 "Lack of rigorous pre-market endurance testing under tropical conditions."
             ])

    # ==========================================
    # SLIDE 6: Financial & Capital Challenges
    # ==========================================
    slide6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide6)
    add_header(slide6, "5. Financial Anatomy: The Cost of Hardware")
    add_card(slide6, Inches(0.8), Inches(2.1), Inches(5.6), Inches(4.8), 
             "Capital Intensity", 
             [
                 "Automotive manufacturing is one of the most capital-intensive industries globally.",
                 "Underestimated cash flow requirements for working capital and inventory.",
                 "Difficulty in securing long-term institutional bank financing in West Africa."
             ])
    add_card(slide6, Inches(6.8), Inches(2.1), Inches(5.7), Inches(4.8), 
             "Unit Economics Failure", 
             [
                 "Low production volume prevented economies of scale (high per-unit fixed cost).",
                 "Inability to compete on price against mass-produced Asian and European imports.",
                 "Rapid cash burn leading to operational paralysis before reaching break-even."
             ])

    # ==========================================
    # SLIDE 7: Industrial & Supply Chain Hurdles
    # ==========================================
    slide7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide7)
    add_header(slide7, "6. Supply Chain & Ecosystem Gaps")
    add_card(slide7, Inches(0.8), Inches(2.1), Inches(5.6), Inches(4.8), 
             "Supply Chain Fragility", 
             [
                 "Heavy reliance on imported components created severe logistics bottlenecks.",
                 "Customs delays, shipping cost volatility, and foreign exchange exposure.",
                 "Absence of local tier-1 and tier-2 automotive parts suppliers."
             ])
    add_card(slide7, Inches(6.8), Inches(2.1), Inches(5.7), Inches(4.8), 
             "Ecosystem Isolation", 
             [
                 "Building a car requires a dense web of specialized sub-contractors.",
                 "Lack of skilled local tooling technicians, specialized metallurgists, and labs.",
                 "Operating in an isolated vacuum without synergistic industrial clusters."
             ])

    # ==========================================
    # SLIDE 8: Market Realities & Competition
    # ==========================================
    slide8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide8)
    add_header(slide8, "7. Market Realities & Consumer Trust")
    add_card(slide8, Inches(0.8), Inches(2.1), Inches(5.6), Inches(4.8), 
             "Fierce Competition", 
             [
                 "Domination of cheap, reliable, second-hand vehicles ('Occasions').",
                 "Established Japanese and Korean brands offering bulletproof reliability and parts availability.",
                 "Consumer risk aversion: hesitation to invest savings in an unproven brand."
             ])
    add_card(slide8, Inches(6.8), Inches(2.1), Inches(5.7), Inches(4.8), 
             "After-Sales & Maintenance Deficit", 
             [
                 "Lack of a widespread garage network, spare parts availability, and warranty backing.",
                 "Resale value uncertainty discouraged pragmatic fleet buyers and taxi syndicates.",
                 "Brand trust eroded rapidly upon encountering initial technical glitches."
             ])

    # ==========================================
    # SLIDE 9: Regulatory & Institutional Gaps
    # ==========================================
    slide9 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide9)
    add_header(slide9, "8. Regulatory & Institutional Framework")
    add_card(slide9, Inches(0.8), Inches(2.1), Inches(5.6), Inches(4.8), 
             "Inadequate Incentives", 
             [
                 "Absence of protective tariff barriers or robust local content laws for infant industries.",
                 "Tax regimes ill-suited for heavy manufacturing investments.",
                 "Limited public procurement support (e.g., state fleets adopting local models)."
             ])
    add_card(slide9, Inches(6.8), Inches(2.1), Inches(5.7), Inches(4.8), 
             "Infrastructure Deficits", 
             [
                 "High electricity costs and unstable grid supply affecting assembly line efficiency.",
                 "Logistics infrastructure limitations for heavy freight transport.",
                 "Weak intellectual property and standardization testing frameworks."
             ])

    # ==========================================
    # SLIDE 10: Post-Mortem Analysis
    # ==========================================
    slide10 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide10)
    add_header(slide10, "9. Post-Mortem: Engineering vs. Market Reality")
    add_card(slide10, Inches(0.8), Inches(2.1), Inches(5.6), Inches(4.8), 
             "Engineering Fallacy", 
             [
                 "Believing that building a working prototype equates to a viable commercial product.",
                 "Underestimating the complexity of quality assurance, ISO standards, and safety certification.",
                 "Focusing purely on product mechanics while ignoring distribution economics."
             ])
    add_card(slide10, Inches(6.8), Inches(2.1), Inches(5.7), Inches(4.8), 
             "Strategic Missteps", 
             [
                 "Premature scaling without validating product-market fit at scale.",
                 "Inadequate risk management against supply chain shocks and currency fluctuations.",
                 "Failure to secure strategic joint ventures with established global automotive giants."
             ])

    # ==========================================
    # SLIDE 11: Key Lessons for Future Industrialists
    # ==========================================
    slide11 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide11)
    add_header(slide11, "10. Key Lessons for African Industrialization")
    add_card(slide11, Inches(0.8), Inches(2.1), Inches(5.6), Inches(4.8), 
             "Ecosystem First", 
             [
                 "Never build an assembly plant without first fostering a supplier and parts ecosystem.",
                 "Start with simpler mechanical or software components before tackling complex vehicle assembly.",
                 "Invest heavily in technical vocational training and specialized engineering schools."
             ])
    add_card(slide11, Inches(6.8), Inches(2.1), Inches(5.7), Inches(4.8), 
             "Pragmatic Positioning", 
             [
                 "Target specific niches first (e.g., electric tricycles, agricultural tractors, utility vans).",
                 "Build robust after-sales service and spare parts availability from day one.",
                 "Align product pricing strictly with local purchasing power realities."
             ])

    # ==========================================
    # SLIDE 12: The "Fail Fast, Learn Faster" Paradigm
    # ==========================================
    slide12 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide12)
    add_header(slide12, "11. The 'Fail Fast, Learn Faster' Paradigm")
    add_card(slide12, Inches(0.8), Inches(2.1), Inches(5.6), Inches(4.8), 
             "A Necessary Pioneer", 
             [
                 "Kpandji acted as the 'sacrificial lamb' of West African automotive ambition.",
                 "It exposed the hidden friction points of local manufacturing that textbooks omit.",
                 "Transformed theoretical debates into hard, empirical industrial data."
             ])
    add_card(slide12, Inches(6.8), Inches(2.1), Inches(5.7), Inches(4.8), 
             "An Asset for Engineers", 
             [
                 "For software and hardware engineers at IIT, Kpandji is a masterclass in risk analysis.",
                 "Teaches us to evaluate system dependencies, financial burn rates, and market friction.",
                 "Lays the empirical groundwork for future successful technological ventures."
             ])

    # ==========================================
    # SLIDE 13: Conclusion
    # ==========================================
    slide13 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide13)
    add_header(slide13, "12. Conclusion: The Future is Still to Be Built")
    add_card(slide13, Inches(0.8), Inches(2.1), Inches(5.6), Inches(4.8), 
             "Beyond the Failure", 
             [
                 "Kpandji's failure is not the end of Ivorian industrial ambition; it is its true starting point.",
                 "Future projects will benefit from Kpandji's trial-and-error legacy.",
                 "The rise of African tech and engineering talent ensures better-prepared successors."
             ])
    add_card(slide13, Inches(6.8), Inches(2.1), Inches(5.7), Inches(4.8), 
             "Core Takeaway", 
             [
                 "« Kpandji's failure is not the end of the Ivorian automotive dream; it is the harsh apprenticeship without which no great edifice can rise. »",
                 "Tomorrow, it is upon stronger, data-driven foundations that our generation will write the future."
             ])

    # ==========================================
    # SLIDE 14: Q&A / Thank You
    # ==========================================
    slide14 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide14)
    
    tb14 = slide14.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.33), Inches(3.5))
    tf14 = tb14.text_frame
    tf14.word_wrap = True

    p = tf14.paragraphs[0]
    p.text = "Thank You for Your Attention!"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(12)

    p_sub = tf14.add_paragraph()
    p_sub.text = "Questions & Discussion"
    p_sub.font.size = Pt(22)
    p_sub.font.color.rgb = EMERALD
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.space_after = Pt(24)

    p_meta = tf14.add_paragraph()
    p_meta.text = "Prepared by Elie Pethuel Ehoussou | IIT L3 Software Engineering\nRepository: Kpandji-Failure-Analysis"
    p_meta.font.size = Pt(14)
    p_meta.font.color.rgb = TEXT_MUTED
    p_meta.alignment = PP_ALIGN.CENTER

    # Save presentation
    output_path = "kpandji_presentation.pptx"
    prs.save(output_path)
    print(f"Presentation successfully saved to {output_path}")

if __name__ == "__main__":
    create_presentation()
